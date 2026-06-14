# -*- coding: utf-8 -*-
"""模板占位符解析：从用户上传的 docx/xlsx 中提取 {{字段}} 占位符，自动生成字段列表。

占位符约定：
    {{树木编号}}            -> key=树木编号, label=树木编号
    {{tree_no|树木编号}}    -> key=tree_no,  label=树木编号
保留占位符(自动忽略，导出时系统填充)：{{_title}} {{_project}} {{_code}} {{_date}}
"""
import re

from docx import Document
from openpyxl import load_workbook

PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
RESERVED = {"_title", "_project", "_code", "_date"}


def parse_token(token):
    """'tree_no|树木编号' -> ('tree_no','树木编号')；'树木编号' -> ('树木编号','树木编号')"""
    token = token.strip()
    if "|" in token:
        key, label = token.split("|", 1)
        return key.strip(), label.strip()
    return token, token


def _fields_from_texts(texts):
    """从一组文本中按出现顺序提取去重的字段。"""
    fields, seen = [], set()
    for t in texts:
        for m in PLACEHOLDER_RE.findall(t or ""):
            key, label = parse_token(m)
            if key in RESERVED or key in seen:
                continue
            seen.add(key)
            fields.append({"key": key, "label": label, "type": "text", "required": False})
    return fields


def _docx_texts(path):
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return texts


def _xlsx_texts(path):
    wb = load_workbook(path, data_only=False)
    texts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    texts.append(cell.value)
    return texts


def extract_fields(path):
    """根据扩展名解析，返回字段列表。"""
    lower = path.lower()
    if lower.endswith(".docx"):
        return _fields_from_texts(_docx_texts(path))
    if lower.endswith(".xlsx"):
        return _fields_from_texts(_xlsx_texts(path))
    raise ValueError("仅支持 .docx 或 .xlsx 模板")
