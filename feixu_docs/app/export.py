# -*- coding: utf-8 -*-
"""导出引擎：把填报记录套用用户模板生成 Word/Excel。
模式：
1) docx 模板：替换 {{key}} / {{key|label}} 占位符
2) xlsx 模板：替换单元格中的占位符
3) 无模板：按字段自动生成两列 Word 表
保留占位符：_title / _project / _code / _date 由系统自动填充。
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import load_workbook

from .constants import PHASE_MAP, PROCESS_MAP
from .template_parse import PLACEHOLDER_RE, parse_token


def _build_mapping(record, template):
    m = dict(record.data)
    m.update({
        "_title": record.title or template.title,
        "_project": record.project.name if record.project else "",
        "_code": template.code or "",
        "_date": record.created_at.strftime("%Y-%m-%d"),
    })
    return m


def _sub_text(text, mapping):
    def repl(match):
        key, _ = parse_token(match.group(1))
        return str(mapping.get(key, ""))
    return PLACEHOLDER_RE.sub(repl, text)


def _fill_docx(template_file, mapping, out_path):
    doc = Document(template_file)

    def fill_paragraph(p):
        full = "".join(run.text for run in p.runs)
        if "{{" not in full:
            return
        new = _sub_text(full, mapping)
        for run in p.runs:
            run.text = ""
        if p.runs:
            p.runs[0].text = new
        else:
            p.add_run(new)

    for p in doc.paragraphs:
        fill_paragraph(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    fill_paragraph(p)
    doc.save(out_path)


def _fill_xlsx(template_file, mapping, out_path):
    wb = load_workbook(template_file)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = _sub_text(cell.value, mapping)
    wb.save(out_path)


def _auto_docx(record, template, out_path):
    data = record.data
    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(10.5)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(template.title)
    r.bold = True
    r.font.size = Pt(15)

    meta = doc.add_paragraph()
    meta.add_run("资料编号：%s    阶段：%s    工艺：%s" % (
        template.code or "—", PHASE_MAP.get(template.phase, ""),
        PROCESS_MAP.get(template.process, ""))).font.size = Pt(9)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for f in template.fields:
        cells = table.add_row().cells
        cells[0].text = f.get("label", f.get("key"))
        cells[1].text = str(data.get(f.get("key"), ""))

    doc.add_paragraph()
    sign = doc.add_table(rows=1, cols=3)
    sign.style = "Table Grid"
    sc = sign.rows[0].cells
    sc[0].text = "施工单位（盖章）："
    sc[1].text = "监理单位："
    sc[2].text = "日期：%s" % record.created_at.strftime("%Y-%m-%d")
    doc.save(out_path)


def export_record(record, template, out_dir, template_file=None):
    """生成导出文件，返回 (文件绝对路径, 扩展名)。"""
    mapping = _build_mapping(record, template)
    # 去掉 Windows 文件名非法字符，保留中文
    safe = re.sub(r'[\\/:*?"<>|]', "_", (template.code or template.title or "record"))
    base = "C-%04d_%s" % (record.id, safe)
    if template_file and os.path.exists(template_file):
        if template_file.lower().endswith(".xlsx"):
            out_path = os.path.join(out_dir, base + ".xlsx")
            _fill_xlsx(template_file, mapping, out_path)
            return out_path, ".xlsx"
        out_path = os.path.join(out_dir, base + ".docx")
        _fill_docx(template_file, mapping, out_path)
        return out_path, ".docx"
    out_path = os.path.join(out_dir, base + ".docx")
    _auto_docx(record, template, out_path)
    return out_path, ".docx"
